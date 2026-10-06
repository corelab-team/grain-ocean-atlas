"""Update one export year from a client workbook without replacing historical years."""
import argparse
import json
from pathlib import Path
import openpyxl
import build_data as bd


def update(source, year, as_of):
    root = Path(bd.ROOT)
    data_path = root / 'data/export.json'
    data = json.loads(data_path.read_text(encoding='utf-8'))
    ref = json.loads((root / 'data/countries_ru.json').read_text(encoding='utf-8'))
    alias = {bd.nrm(a): n for n, r in ref.items() for a in r.get('xlsx', [])}
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = bd.find_sheet(wb, 'страны') or bd.find_sheet(wb, bd.SHEET_COUNTRIES)
    if ws is None:
        raise ValueError('Country sheet is missing')
    header, cols = bd.read_header_years(ws)
    cols = [(c, y) for c, y in cols if y == year]
    if len(cols) != 1 or year not in data['years']:
        raise ValueError('Requested year is missing')
    groups, products, rows, group_totals = bd.parse_countries_sheet(ws, header, cols)
    index = {(data['groups'][p['g']], p['n']): i for i, p in enumerate(data['products'])}
    product_map = {}
    for i, product in enumerate(products):
        key = (groups[product['g']], product['n'])
        if key not in index:
            raise ValueError('Unknown product: ' + str(key))
        product_map[i] = index[key]
    exact = {}
    for name, values in rows.items():
        canonical = alias.get(name, name)
        bucket = exact.setdefault(canonical, {})
        for product, values_by_year in values.items():
            value = values_by_year.get(year, 0)
            if value:
                pi = product_map[product]
                bucket[pi] = bucket.get(pi, 0) + value
    active = {n: v for n, v in exact.items() if v}
    missing = set(active) - set(ref)
    if missing:
        raise ValueError('Missing coordinates: ' + str(sorted(missing)))
    expected = sum(group_totals[year].values())
    actual = sum(sum(v.values()) for v in active.values())
    if abs(expected - actual) > 0.001:
        raise ValueError('Country/group totals differ')
    source_total = next(bd.num(ws.cell(r, cols[0][0]).value) for r in range(header + 1, ws.max_row + 1)
                        if bd.nrm(ws.cell(r, 1).value).lower() == bd.TOTAL_ROW)
    if abs(actual - source_total) > 0.001:
        raise ValueError('Workbook total differs')
    countries = {c['name']: c for c in data['countries']}
    for name in active:
        if name not in countries:
            r = ref[name]
            c = dict(name=name, en=r.get('en', name), iso=r.get('iso'), lat=r['lat'], lon=r['lon'],
                     totals=[0] * len(data['years']), years={})
            if r.get('port'): c['port'] = r['port']
            data['countries'].append(c)
            countries[name] = c
    slot = data['years'].index(year)
    for name, country in countries.items():
        values = active.get(name, {})
        country['totals'][slot] = round(sum(values.values()), 9)
        country['years'].pop(str(year), None)
        if values:
            country['years'][str(year)] = sorted([[i, round(v, 9)] for i, v in values.items()], key=lambda x: -x[1])
    ranked = sorted(((n, sum(v.values())) for n, v in active.items()), key=lambda x: -x[1])
    used_groups = {data['products'][i]['g'] for v in active.values() for i in v}
    data['summary'][str(year)] = dict(total=round(actual, 3), countries=len(active), groups=len(used_groups),
                                    top=[dict(name=n, value=round(v, 3)) for n, v in ranked[:3]])
    data.setdefault('sourcesByYear', {})[str(year)] = dict(file=Path(source).name, asOf=as_of)
    data_path.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    config_path = root / 'config.json'
    config = json.loads(config_path.read_text(encoding='utf-8'))
    config.setdefault('yearNote', {})[str(year)] = 'на ' + as_of
    config_path.write_text(json.dumps(config, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'{year}: {actual:.3f} thousand tonnes, {len(active)} countries, {len(used_groups)} groups')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('--year', type=int, default=2026)
    parser.add_argument('--as-of', required=True)
    args = parser.parse_args()
    update(args.source, args.year, args.as_of)
